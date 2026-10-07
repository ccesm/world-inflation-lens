"""Microsoft's official historical native CapEx workbook. No allocation to AI."""
import re,json,hashlib,datetime,calendar,sys
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

def extract_workbook(source,cache):
 p=Path(cache)/'objects'/source['sha256'][:2]/source['sha256'];b=p.read_bytes()
 if hashlib.sha256(b).hexdigest()!=source['sha256']:raise ValueError('RAW_HASH')
 w=load_workbook(BytesIO(b),data_only=True);s=w['CapEx']
 if 'Capital Expenditures Including Assets Acquired Under Capital Leases' not in s.cell(2,1).value or s.cell(3,1).value!='(In billions)':raise ValueError('WORKBOOK_DEFINITION')
 headers=[r for r in s if any(isinstance(c.value,str) and re.fullmatch(r'Q[1-4]-\d{2}',c.value) for c in r)]
 if len(headers)!=1:raise ValueError('WORKBOOK_HEADERS')
 header=headers[0];data=s[header[0].row+1];rows=[]
 for cell in header:
  match=re.fullmatch(r'Q([1-4])-(\d{2})',str(cell.value))
  if not match:continue
  q,y=int(match[1]),2000+int(match[2]);m={1:9,2:12,3:3,4:6}[q];cy=y-(q<3);end=f'{cy}-{m:02d}-{calendar.monthrange(cy,m)[1]}'
  if not '2019-03-31'<=end<='2026-06-30':continue
  value=data[cell.column-1].value
  if not isinstance(value,(int,float)):raise ValueError('WORKBOOK_VALUE')
  rows.append({'observationId':f"{source['sourceId']}:nativeCapex:{end}",'sourceId':source['sourceId'],'company':'MSFT','metric':'nativeCapex','scope':'CONSOLIDATED','periodStart':f'{cy}-{m-2:02d}-01','periodEnd':end,'periodType':'Q','unit':'USD_MILLIONS','value':round(value*1000,8),'precision':'ROUNDED','evidenceClass':'OBSERVED','definitionVersion':'MICROSOFT_CAPEX_V1','restatementBasis':source['sourceId'],'revisionStatus':'CURRENT_HISTORICAL_WORKBOOK_VINTAGE','nativeLabel':data[0].value,'locator':f'CapEx!{data[cell.column-1].coordinate}; header {cell.coordinate}={cell.value}','comparabilityStatus':'LIMITED_COMPARABILITY'})
 w.close()
 return rows
if __name__=='__main__':
 source=json.load(open(sys.argv[1]));json.dump(extract_workbook(source,sys.argv[2]),open(sys.argv[3],'w'),indent=2)
